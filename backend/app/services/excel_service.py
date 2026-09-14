import io
import logging
from datetime import datetime
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.event import Event
from app.models.order import Order
from app.services.pdf_service import build_checkin_rows, format_french_date

logger = logging.getLogger(__name__)

# Mandatory 8 columns specified for official check-in
COLUMNS_DEF = [
    ("N° Stand", 16, "center"),
    ("Métrage", 12, "center"),
    ("Nom / Prénom", 32, "left"),
    ("Téléphone", 18, "center"),
    ("Statut paiement", 18, "center"),
    ("Présent", 12, "center"),
    ("Pièce d'identité contrôlée", 26, "center"),
    ("N° CNI relevé", 25, "left"),
]


def sanitize_cell_value(val):
    """
    Sanitize string values against CSV/Excel formula injection.
    Prepends a single quote if the string starts with '=', '+', '-', or '@'.
    """
    if isinstance(val, str) and val.startswith(("=", "+", "-", "@")):
        return "'" + val
    return val


def _style_sheet(
    ws,
    event: Event,
    orders: list[Order],
    sort_by: str,
    sheet_title: str,
) -> None:
    """Populate and format a worksheet for the official check-in registry."""
    rows = build_checkin_rows(orders=orders, sort_by=sort_by)
    event_date_formatted = format_french_date(event.start_date)

    # Freeze panes below the header row so rows 1 to 4 remain visible on scroll
    ws.freeze_panes = "A5"

    # Configure print setup: Landscape, A4, fit to 1 page wide
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.fitToPage = True

    # Styling definitions
    title_font = Font(name="Calibri", size=14, bold=True, color="065F46")
    meta_font = Font(name="Calibri", size=10, italic=True, color="475569")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Row 1: Event Title
    ws["A1"] = f"Feuille d'Émargement Officielle — {event.title}"
    ws["A1"].font = title_font
    ws.merge_cells("A1:H1")

    # Row 2: Event Details, Sort Mode & Edition Timestamp
    now = datetime.now()
    timestamp_str = f"Édité le {now.day:02d}/{now.month:02d}/{now.year} à {now.hour:02d}h{now.minute:02d}"

    mode_text = (
        "Classement : Par Emplacement (N° Stand)"
        if sort_by == "spot"
        else "Classement : Alphabétique (Nom exposant)"
    )
    location = event.location_address or "Non spécifié"
    total_meters = sum(r["linear_meters"] for r in rows)
    meters_str = f"{total_meters:.2f}".replace(".", ",") + " m"

    ws["A2"] = (
        f"Date : {event_date_formatted} | Lieu : {location} | "
        f"{mode_text} | Total : {len(rows)} entrées ({meters_str}) | {timestamp_str}"
    )
    ws["A2"].font = meta_font
    ws.merge_cells("A2:H2")

    # Row 3 is an empty separator row

    # Row 4: Column Headers
    header_row_idx = 4
    for col_idx, (col_name, _, align) in enumerate(COLUMNS_DEF, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row_idx].height = 24

    # Row 5+: Data rows
    if not rows:
        empty_cell = ws.cell(row=5, column=1, value="Aucun exposant inscrit pour cet événement.")
        empty_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells("A5:H5")
        ws.row_dimensions[5].height = 26
        for col_idx in range(1, 9):
            ws.cell(row=5, column=col_idx).border = thin_border
    else:
        for r_idx, row_data in enumerate(rows, start=5):
            is_alt = (r_idx % 2 == 0)
            ws.row_dimensions[r_idx].height = 20

            meters_numeric = float(row_data["linear_meters"])
            values = [
                sanitize_cell_value(row_data["spot_label"]),
                meters_numeric,
                sanitize_cell_value(row_data["full_name"]),
                sanitize_cell_value(row_data["phone"]),
                sanitize_cell_value(row_data["payment_status"]),
                "",  # Checkbox "Présent"
                "",  # Checkbox "Pièce d'identité contrôlée"
                "",  # Dotted pen entry for CNI
            ]

            for c_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                align = COLUMNS_DEF[c_idx - 1][2]
                cell.alignment = Alignment(horizontal=align, vertical="center")
                cell.border = thin_border

                if is_alt:
                    cell.fill = alt_row_fill

                if c_idx == 1:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="065F46")
                elif c_idx == 2:
                    # Numeric linear meters with standard unit formatting
                    cell.number_format = '0.00 "m"'
                    cell.font = Font(name="Calibri", size=10, color="1E293B")
                elif c_idx == 3:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
                else:
                    cell.font = Font(name="Calibri", size=10, color="1E293B")

    # Set column widths
    for c_idx, (_, width, _) in enumerate(COLUMNS_DEF, start=1):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = width


def generate_checkin_xlsx(
    event: Event,
    orders: list[Order],
    sort_by: str = "spot",
) -> bytes:
    """
    Generate an official check-in sheet workbook (.xlsx) with styled headers,
    freeze panes, landscape A4 print setup, adjusted column widths, and pen-ready check-in columns.
    Creates two worksheets:
      1. 'Par Emplacement' (sorted naturally by spot number)
      2. 'Par Nom (Alphabétique)' (sorted alphabetically by exhibitor name)
    The active sheet is set based on sort_by.
    Returns raw binary XLSX bytes directly from memory.
    """
    wb = openpyxl.Workbook()

    # Sheet 1: Spot order
    ws_spot = wb.active
    ws_spot.title = "Par Emplacement"
    _style_sheet(
        ws=ws_spot,
        event=event,
        orders=orders,
        sort_by="spot",
        sheet_title="Par Emplacement",
    )

    # Sheet 2: Alpha order
    ws_alpha = wb.create_sheet(title="Par Nom (Alphabétique)")
    _style_sheet(
        ws=ws_alpha,
        event=event,
        orders=orders,
        sort_by="alpha",
        sheet_title="Par Nom (Alphabétique)",
    )

    # Set active sheet according to sort_by
    if sort_by == "alpha":
        wb.active = ws_alpha
    else:
        wb.active = ws_spot

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
