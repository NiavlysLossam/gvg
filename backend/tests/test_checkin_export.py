import io
import uuid
import openpyxl
from datetime import datetime, timezone
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.event import Event
from app.models.spot import Spot
from app.models.order import Order, BookingItem
from app.services import pdf_service, excel_service
from app.services.pdf_service import (
    natural_sort_key,
    french_alpha_sort_key,
    build_checkin_rows,
    generate_checkin_pdf,
    pdf_jinja_env,
)
from app.services.excel_service import generate_checkin_xlsx, COLUMNS_DEF


# ==============================================================================
# Helpers & Fixtures
# ==============================================================================

def create_event_with_multiple_orders(db: Session) -> tuple[Event, list[Order]]:
    """Helper creating an event with confirmed and non-confirmed orders."""
    event = Event(
        title="Grand Vide-Greniers du Printemps",
        slug=f"vg-printemps-{uuid.uuid4().hex[:6]}",
        price_per_meter_cents=500,
        map_type="geographic",
        start_date=datetime(2026, 5, 24, 7, 0, tzinfo=timezone.utc),
        end_date=datetime(2026, 5, 24, 18, 0, tzinfo=timezone.utc),
        setup_start_time="06:00",
        setup_end_time="07:30",
        location_address="Place de la Mairie, 44000 Nantes",
        organizer_email="organisateur@nantes-vg.fr",
        status="published",
    )
    db.add(event)
    db.flush()

    # Spots: A-1, A-2, A-10, B-1
    spot_a1 = Spot(
        event_id=event.id,
        label="A-1",
        linear_meters=2.0,
        price_cents=1000,
        status="reserved",
        geom="SRID=-1;POLYGON((0 0, 0 2, 2 2, 2 0, 0 0))",
    )
    spot_a2 = Spot(
        event_id=event.id,
        label="A-2",
        linear_meters=2.5,
        price_cents=1250,
        status="reserved",
        geom="SRID=-1;POLYGON((2 0, 2 2, 4.5 2, 4.5 0, 2 0))",
    )
    spot_a10 = Spot(
        event_id=event.id,
        label="A-10",
        linear_meters=3.0,
        price_cents=1500,
        status="reserved",
        geom="SRID=-1;POLYGON((5 0, 5 2, 8 2, 8 0, 5 0))",
    )
    spot_b1 = Spot(
        event_id=event.id,
        label="B-1",
        linear_meters=4.0,
        price_cents=2000,
        status="reserved",
        geom="SRID=-1;POLYGON((0 3, 0 5, 4 5, 4 3, 0 3))",
    )
    db.add_all([spot_a1, spot_a2, spot_a10, spot_b1])
    db.flush()

    # Order 1: François Dupont -> Spot A-10 (confirmed)
    order1 = Order(
        event_id=event.id,
        order_number=f"GVG-2026-ORD1",
        first_name="François",
        last_name="Dupont",
        email="francois.dupont@example.com",
        phone="06 11 22 33 44",
        status="confirmed",
        payment_method="stripe",
        total_price_cents=1500,
    )
    db.add(order1)
    db.flush()
    db.add(BookingItem(order_id=order1.id, spot_id=spot_a10.id, price_cents=1500))

    # Order 2: Éric Benoît -> Spot A-1 and A-2 (confirmed, multi-spot, accent in first name)
    order2 = Order(
        event_id=event.id,
        order_number=f"GVG-2026-ORD2",
        first_name="Éric",
        last_name="Benoît",
        email="eric.benoit@example.com",
        phone="06 55 66 77 88",
        status="confirmed",
        payment_method="cash",
        total_price_cents=2250,
    )
    db.add(order2)
    db.flush()
    db.add(BookingItem(order_id=order2.id, spot_id=spot_a1.id, price_cents=1000))
    db.add(BookingItem(order_id=order2.id, spot_id=spot_a2.id, price_cents=1250))

    # Order 3: Alain Martin -> Spot B-1 (confirmed, paid check)
    order3 = Order(
        event_id=event.id,
        order_number=f"GVG-2026-ORD3",
        first_name="Alain",
        last_name="Martin",
        email="alain.martin@example.com",
        phone="06 99 88 77 66",
        status="confirmed",
        payment_method="check",
        total_price_cents=2000,
    )
    db.add(order3)
    db.flush()
    db.add(BookingItem(order_id=order3.id, spot_id=spot_b1.id, price_cents=2000))

    # Order 4: Cancelled order (must be excluded from check-in)
    order_cancelled = Order(
        event_id=event.id,
        order_number=f"GVG-2026-CANC",
        first_name="Sophie",
        last_name="Annulée",
        phone="06 00 00 00 01",
        status="cancelled",
        payment_method="stripe",
        total_price_cents=1000,
    )
    db.add(order_cancelled)

    # Order 5: Refunded order (must be excluded from check-in)
    order_refunded = Order(
        event_id=event.id,
        order_number=f"GVG-2026-REFU",
        first_name="Robert",
        last_name="Remboursé",
        phone="06 00 00 00 02",
        status="refunded",
        payment_method="stripe",
        total_price_cents=1000,
    )
    db.add(order_refunded)

    # Order 6: Pending order (must be excluded from check-in)
    order_pending = Order(
        event_id=event.id,
        order_number=f"GVG-2026-PEND",
        first_name="Paul",
        last_name="EnAttente",
        phone="06 00 00 00 03",
        status="pending",
        payment_method="stripe",
        total_price_cents=1000,
    )
    db.add(order_pending)

    db.commit()
    db.refresh(event)

    return event, [order1, order2, order3, order_cancelled, order_refunded, order_pending]


# ==============================================================================
# Unit Tests : Natural & French Sorting Helpers
# ==============================================================================

def test_natural_sort_key_alphanumeric():
    """Verify natural alphanumeric sorting (A-1, A-2, A-10, B-1) and null guard."""
    assert natural_sort_key(None) == []
    assert natural_sort_key("") == []
    labels = ["B-1", "A-10", "A-2", "A-1"]
    sorted_labels = sorted(labels, key=natural_sort_key)
    assert sorted_labels == ["A-1", "A-2", "A-10", "B-1"]


def test_french_alpha_sort_key_accents_and_case():
    """Verify French accent-insensitive and case-insensitive collation (e.g. Éric before François)."""
    names = ["François", "Éric", "alain", "Benoît", "cécile", "Cecile"]
    sorted_names = sorted(names, key=french_alpha_sort_key)
    # Normalized: alain (a), benoit (b), cecile (c), eric (e), francois (f)
    assert sorted_names[0] == "alain"
    assert sorted_names[1] == "Benoît"
    # Éric comes before François
    eric_idx = sorted_names.index("Éric")
    francois_idx = sorted_names.index("François")
    assert eric_idx < francois_idx


# ==============================================================================
# Unit Tests : build_checkin_rows
# ==============================================================================

def test_build_checkin_rows_spot_mode(db_session: Session):
    """In spot mode, rows are expanded per spot and sorted naturally by stall number."""
    event, orders = create_event_with_multiple_orders(db_session)

    rows = build_checkin_rows(orders=orders, sort_by="spot")

    # Only 3 confirmed orders -> 4 spots total (A-1, A-2 from order 2, A-10 from order 1, B-1 from order 3)
    assert len(rows) == 4

    spot_labels = [r["spot_label"] for r in rows]
    assert spot_labels == ["A-1", "A-2", "A-10", "B-1"]

    # Verify order 2 (Éric Benoît) appears on both A-1 and A-2
    assert rows[0]["full_name"] == "Éric Benoît"
    assert rows[0]["spot_label"] == "A-1"
    assert rows[0]["linear_meters_str"] == "2,00 m"
    assert rows[0]["payment_status"] == "Payé (Espèces)"

    assert rows[1]["full_name"] == "Éric Benoît"
    assert rows[1]["spot_label"] == "A-2"
    assert rows[1]["linear_meters_str"] == "2,50 m"

    # Verify order 1 (François Dupont) on A-10
    assert rows[2]["full_name"] == "François Dupont"
    assert rows[2]["spot_label"] == "A-10"
    assert rows[2]["payment_status"] == "Payé (CB)"

    # Verify order 3 (Alain Martin) on B-1
    assert rows[3]["full_name"] == "Alain Martin"
    assert rows[3]["spot_label"] == "B-1"
    assert rows[3]["payment_status"] == "Payé (Chèque)"

    # Verify unconfirmed orders are absent
    names = [r["full_name"] for r in rows]
    assert "Sophie Annulée" not in names
    assert "Robert Remboursé" not in names
    assert "Paul EnAttente" not in names


def test_build_checkin_rows_alpha_mode(db_session: Session):
    """In alpha mode, multi-spot orders are grouped and sorted by exhibitor name."""
    event, orders = create_event_with_multiple_orders(db_session)

    rows = build_checkin_rows(orders=orders, sort_by="alpha")

    # Exactly 3 confirmed orders -> 3 rows
    assert len(rows) == 3

    # Expected order by last name: Benoît (Éric), Dupont (François), Martin (Alain)
    names = [r["full_name"] for r in rows]
    assert names == ["Éric Benoît", "François Dupont", "Alain Martin"]

    # Verify multi-spot grouping on Éric Benoît: A-1, A-2 (4.50 m)
    benoit_row = rows[0]
    assert benoit_row["spot_label"] == "A-1, A-2"
    assert benoit_row["linear_meters_str"] == "4,50 m"
    assert benoit_row["linear_meters"] == 4.5


def test_build_checkin_rows_empty_orders():
    """Empty or unconfirmed-only list produces an empty row list without error."""
    rows = build_checkin_rows(orders=[], sort_by="spot")
    assert rows == []

    # Unconfirmed orders
    cancelled_order = Order(first_name="Test", last_name="Cancelled", status="cancelled")
    rows_unconfirmed = build_checkin_rows(orders=[cancelled_order], sort_by="alpha")
    assert rows_unconfirmed == []


# ==============================================================================
# Unit Tests : PDF Generation (generate_checkin_pdf & Template)
# ==============================================================================

def test_generate_checkin_pdf_spot_mode(db_session: Session):
    """PDF check-in in spot mode generates valid A4 landscape PDF bytes."""
    event, orders = create_event_with_multiple_orders(db_session)

    pdf_bytes = generate_checkin_pdf(event=event, orders=orders, sort_by="spot")

    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 5000


def test_generate_checkin_pdf_alpha_mode(db_session: Session):
    """PDF check-in in alpha mode generates valid PDF bytes."""
    event, orders = create_event_with_multiple_orders(db_session)

    pdf_bytes = generate_checkin_pdf(event=event, orders=orders, sort_by="alpha")

    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 5000


def test_checkin_html_template_contains_mandatory_headers_and_elements(db_session: Session):
    """Verify the rendered checkin.html contains all required table headers and check-in elements."""
    event, orders = create_event_with_multiple_orders(db_session)
    rows = build_checkin_rows(orders=orders, sort_by="spot")

    template = pdf_jinja_env.get_template("checkin.html")
    rendered = template.render(
        event=event,
        rows=rows,
        sort_by="spot",
        sort_label="Par Emplacement",
        mode_label="Classement : Par Emplacement (N° Stand)",
        event_date_formatted="24 mai 2026",
        total_rows=len(rows),
        total_linear_meters_formatted="11,50 m",
        generation_date="14/09/2026 à 21h00",
    )

    # Mandatory 8 columns
    assert "N° Stand" in rendered
    assert "Métrage" in rendered
    assert "Nom / Prénom" in rendered
    assert "Téléphone" in rendered
    assert "Statut paiement" in rendered
    assert "Présent" in rendered
    assert "Pièce d'identité contrôlée" in rendered
    assert "N° CNI relevé" in rendered

    # CSS requirements for A4 landscape, repeated thead and page numbering
    assert "size: A4 landscape" in rendered
    assert "thead {" in rendered
    assert "display: table-header-group" in rendered
    assert "counter(page)" in rendered

    # Running headers on every printed page
    assert "@top-left {" in rendered
    assert event.title in rendered
    assert "@top-right {" in rendered
    assert "Émargement officiel — Par Emplacement" in rendered

    # Word wrapping for exhibitor names (no truncation on paper)
    assert "word-break: break-word" in rendered
    assert "white-space: normal" in rendered

    # Data elements
    assert "A-1" in rendered
    assert "A-2" in rendered
    assert "A-10" in rendered
    assert "B-1" in rendered
    assert "Éric Benoît" in rendered
    assert "François Dupont" in rendered
    assert "Alain Martin" in rendered



def test_generate_checkin_pdf_empty_event():
    """Event with 0 confirmed orders produces a valid PDF indicating 'Aucun exposant inscrit'."""
    event = Event(
        title="Brocante Vide",
        slug="brocante-vide",
        start_date=datetime(2026, 8, 15, 8, 0),
        price_per_meter_cents=500,
    )

    pdf_bytes = generate_checkin_pdf(event=event, orders=[], sort_by="spot")
    assert pdf_bytes.startswith(b"%PDF-")

    # Check rendered template
    template = pdf_jinja_env.get_template("checkin.html")
    rendered = template.render(
        event=event,
        rows=[],
        sort_by="spot",
        mode_label="Classement : Par Emplacement (N° Stand)",
        event_date_formatted="15 août 2026",
        total_rows=0,
        total_linear_meters_formatted="0,00 m",
        generation_date="14/09/2026 à 21h00",
    )
    assert "Aucun exposant inscrit pour cet événement." in rendered


def test_generate_checkin_pdf_missing_optional_fields():
    """PDF generation succeeds even when exhibitor fields (phone, email) are None."""
    event = Event(
        title="Brocante Spéciale",
        slug="brocante-speciale",
        start_date=datetime(2026, 9, 1, 8, 0),
        price_per_meter_cents=500,
    )
    spot = Spot(label="S-01", linear_meters=2.0, price_cents=1000)
    order = Order(
        first_name="Jean",
        last_name="SansTéléphone",
        phone=None,
        email=None,
        status="confirmed",
        payment_method=None,
        items=[BookingItem(spot=spot, price_cents=1000)],
    )

    pdf_bytes = generate_checkin_pdf(event=event, orders=[order], sort_by="spot")
    assert pdf_bytes.startswith(b"%PDF-")


# ==============================================================================
# Unit Tests : Excel Workbook Generation (generate_checkin_xlsx)
# ==============================================================================

def test_generate_checkin_xlsx_structure_and_columns(db_session: Session):
    """Excel export produces a valid XLSX workbook with 8 required columns and 2 worksheets."""
    event, orders = create_event_with_multiple_orders(db_session)

    xlsx_bytes = generate_checkin_xlsx(event=event, orders=orders, sort_by="spot")

    assert isinstance(xlsx_bytes, bytes)
    # XLSX files are ZIP archives starting with PK\x03\x04
    assert xlsx_bytes.startswith(b"PK\x03\x04")

    # Load with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Par Emplacement" in wb.sheetnames
    assert "Par Nom (Alphabétique)" in wb.sheetnames

    # Check Sheet 1 (Spot order)
    ws_spot = wb["Par Emplacement"]
    assert ws_spot["A1"].value.startswith("Feuille d'Émargement Officielle")

    # Freeze panes & Page Setup
    assert ws_spot.freeze_panes == "A5"
    assert ws_spot.page_setup.orientation == ws_spot.ORIENTATION_LANDSCAPE
    assert ws_spot.page_setup.fitToWidth == 1
    assert "Édité le" in ws_spot["A2"].value

    # Verify column headers in row 4
    expected_headers = [col_def[0] for col_def in COLUMNS_DEF]
    actual_headers = [ws_spot.cell(row=4, column=i).value for i in range(1, 9)]
    assert actual_headers == expected_headers

    # Check data rows in ws_spot (sorted A-1, A-2, A-10, B-1)
    spot_col_vals = [ws_spot.cell(row=r, column=1).value for r in range(5, 9)]
    assert spot_col_vals == ["A-1", "A-2", "A-10", "B-1"]

    # Numeric linear meters with standard number format
    meter_val = ws_spot.cell(row=5, column=2).value
    assert isinstance(meter_val, (int, float))
    assert meter_val == 2.0
    assert ws_spot.cell(row=5, column=2).number_format == '0.00 "m"'

    # Check Sheet 2 (Alpha order)
    ws_alpha = wb["Par Nom (Alphabétique)"]
    names_col_vals = [ws_alpha.cell(row=r, column=3).value for r in range(5, 8)]
    assert names_col_vals == ["Éric Benoît", "François Dupont", "Alain Martin"]

    # Multi-spot grouped in ws_alpha row 5
    assert ws_alpha.cell(row=5, column=1).value == "A-1, A-2"
    assert ws_alpha.cell(row=5, column=2).value == 4.5
    assert ws_alpha.cell(row=5, column=2).number_format == '0.00 "m"'


def test_generate_checkin_xlsx_formula_injection_sanitization():
    """Cells starting with =, +, -, @ are sanitized by prepending a single quote."""
    event = Event(
        title="Brocante Sécurité",
        slug="brocante-secu",
        start_date=datetime(2026, 6, 1, 8, 0),
        price_per_meter_cents=500,
    )
    spot = Spot(label="=CMD|' /C calc'!A0", linear_meters=3.0, price_cents=1500)
    order = Order(
        first_name="+33600000000",
        last_name="-ATTACK",
        phone="@special_phone",
        status="confirmed",
        payment_method="stripe",
        items=[BookingItem(spot=spot, price_cents=1500)],
    )

    xlsx_bytes = generate_checkin_xlsx(event=event, orders=[order], sort_by="spot")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Par Emplacement"]

    # Row 5 data cells must be prefixed with single quote
    assert ws.cell(row=5, column=1).value.startswith("'=")
    assert ws.cell(row=5, column=3).value.startswith("'+33600000000")
    assert ws.cell(row=5, column=4).value.startswith("'@special_phone")


def test_generate_checkin_xlsx_alpha_active_sheet(db_session: Session):
    """When sort_by='alpha', the active sheet is set to 'Par Nom (Alphabétique)'."""
    event, orders = create_event_with_multiple_orders(db_session)

    xlsx_bytes = generate_checkin_xlsx(event=event, orders=orders, sort_by="alpha")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.active.title == "Par Nom (Alphabétique)"



def test_generate_checkin_xlsx_empty_event():
    """Event with 0 confirmed orders produces a clean workbook with empty state message."""
    event = Event(
        title="Vide-Grenier Vide",
        slug="vide-grenier-vide",
        start_date=datetime(2026, 7, 10, 8, 0),
        price_per_meter_cents=500,
    )

    xlsx_bytes = generate_checkin_xlsx(event=event, orders=[], sort_by="spot")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.cell(row=5, column=1).value == "Aucun exposant inscrit pour cet événement."


# ==============================================================================
# Integration Tests : API Endpoints
# ==============================================================================

def test_api_download_checkin_pdf_nominal_spot(client: TestClient, db_session: Session):
    """GET /api/v1/events/{id_or_slug}/checkin.pdf?sort_by=spot returns 200 with PDF content."""
    event, _ = create_event_with_multiple_orders(db_session)

    res = client.get(f"/api/v1/events/{event.slug}/checkin.pdf?sort_by=spot")

    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.headers["cache-control"] == "private, no-store, must-revalidate"
    assert f'filename="emargement_{event.slug}_spot.pdf"' in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")


def test_api_download_checkin_pdf_nominal_alpha(client: TestClient, db_session: Session):
    """GET /api/v1/events/{id_or_slug}/checkin.pdf?sort_by=alpha returns 200 with PDF content."""
    event, _ = create_event_with_multiple_orders(db_session)

    res = client.get(f"/api/v1/events/{event.id}/checkin.pdf?sort_by=alpha")

    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert f'filename="emargement_{event.slug}_alpha.pdf"' in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")


def test_api_download_checkin_pdf_invalid_sort_by(client: TestClient, db_session: Session):
    """GET /api/v1/events/{id_or_slug}/checkin.pdf?sort_by=invalid returns 400 Bad Request."""
    event, _ = create_event_with_multiple_orders(db_session)

    res = client.get(f"/api/v1/events/{event.slug}/checkin.pdf?sort_by=invalid_sort")

    assert res.status_code == 400
    assert "sort_by" in res.json()["detail"]


def test_api_download_checkin_xlsx_nominal(client: TestClient, db_session: Session):
    """GET /api/v1/events/{id_or_slug}/checkin.xlsx returns 200 with Excel MIME type and bytes."""
    event, _ = create_event_with_multiple_orders(db_session)

    res = client.get(f"/api/v1/events/{event.slug}/checkin.xlsx")

    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res.headers["cache-control"] == "private, no-store, must-revalidate"
    assert f'filename="emargement_{event.slug}.xlsx"' in res.headers["content-disposition"]
    assert res.content.startswith(b"PK\x03\x04")


def test_api_download_checkin_not_found(client: TestClient):
    """Endpoints return 404 for unknown event ID or slug."""
    fake_slug = "evenement-fantome-12345"
    res_pdf = client.get(f"/api/v1/events/{fake_slug}/checkin.pdf")
    assert res_pdf.status_code == 404

    res_xlsx = client.get(f"/api/v1/events/{fake_slug}/checkin.xlsx")
    assert res_xlsx.status_code == 404

